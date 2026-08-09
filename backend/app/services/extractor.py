from __future__ import annotations

import json
import logging
import re
from pathlib import Path

import pdfplumber
from docx import Document as DocxDocument
from openpyxl import load_workbook

from app.services.ocr import extract_document, ocr_runtime_status
from app.services.ocr.types import ExtractionPhase

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".xlsx",
    ".xls",
    ".csv",
    ".txt",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".bmp",
    ".webp",
    ".dwg",
    ".dxf",
    ".rvt",
    ".rfa",
    ".rte",
    ".rft",
    ".ifc",
    ".x81",
    ".x82",
    ".x83",
    ".x84",
    ".x85",
    ".x86",
    ".d81",
    ".d82",
    ".d83",
    ".d84",
    ".d85",
    ".d86",
}

_GAEB_EXTENSIONS = {
    ".x81",
    ".x82",
    ".x83",
    ".x84",
    ".x85",
    ".x86",
    ".d81",
    ".d82",
    ".d83",
    ".d84",
    ".d85",
    ".d86",
}

_MAX_DOCX_IMAGES = 12
_MIN_NATIVE_TEXT_CHARS = 40
_MIN_PAGE_TEXT_CHARS = 20


def extract_text_from_file(
    path: Path,
    *,
    allow_ocr: bool = True,
    treat_as_drawing: bool = False,
) -> str:
    """
    Backward-compatible facade used by analyzer/upload.

    PDFs go through the modular DocumentPipeline (detect → OCR → merge).
    """
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            result = extract_document(
                path, allow_ocr=allow_ocr, treat_as_drawing=treat_as_drawing
            )
            if result.merged_text:
                return result.merged_text
            if result.error:
                return f"[EXTRACT_ERROR] {path.name}: {result.error}"
            return f"[OCR_EMPTY] {path.name}\nNo extractable text found."
        if suffix == ".docx":
            return _extract_docx(path, allow_ocr=allow_ocr)
        if suffix in {".xlsx", ".xls"}:
            return _extract_excel(path)
        if suffix in {".txt", ".csv"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
            if not allow_ocr:
                return (
                    f"[OCR_SKIPPED] {path.name}\n"
                    "OCR skipped during bulk upload. Use re-extract if text is needed."
                )
            return _extract_image(path)
        if suffix in {".dwg", ".dxf", ".rvt", ".rfa", ".rte", ".rft"}:
            from app.services.cad_extractor import extract_cad_text

            return extract_cad_text(path)
        if suffix == ".ifc":
            from app.services.ifc_extractor import extract_ifc

            return extract_ifc(path).merged_text
        if suffix in _GAEB_EXTENSIONS:
            from app.services.gaeb_extractor import extract_gaeb

            return extract_gaeb(path).merged_text
        if suffix == ".doc":
            return (
                f"[BINARY_OR_IMAGE_FILE] {path.name}\n"
                "Text extraction for this format is limited in MVP. "
                "Convert .doc to .docx or upload a searchable PDF."
            )
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as exc:  # noqa: BLE001
        return f"[EXTRACT_ERROR] {path.name}: {exc}"


def extract_document_full(
    path: Path,
    *,
    allow_ocr: bool = True,
    treat_as_drawing: bool = False,
    max_pages: int | None = None,
    on_progress=None,
):
    """Full structured extraction (pages JSON + merged text + optional structured meta)."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        kwargs = {
            "allow_ocr": allow_ocr,
            "treat_as_drawing": treat_as_drawing,
            "on_progress": on_progress,
        }
        if max_pages is not None:
            kwargs["max_pages"] = max_pages
        return extract_document(path, **kwargs)
    if suffix == ".ifc":
        from app.services.ifc_extractor import extract_ifc

        return extract_ifc(path).to_document_result()
    if suffix in _GAEB_EXTENSIONS:
        from app.services.gaeb_extractor import extract_gaeb

        return extract_gaeb(path).to_document_result()

    # Non-PDF: wrap flat extract into a minimal result-like dict via pipeline types
    from app.services.ocr.types import DocumentExtractionResult, OcrPageResult

    text = extract_text_from_file(path, allow_ocr=allow_ocr, treat_as_drawing=treat_as_drawing)
    ok = has_usable_text(text)
    method = None
    if suffix == ".dxf":
        method = "dxf_native"
    elif suffix == ".dwg":
        method = "dwg_harvest"
    elif suffix in {".rvt", ".rfa", ".rte", ".rft"}:
        method = "rvt_harvest"
    page = OcrPageResult(
        page=1,
        text=text if ok else "",
        confidence=90.0 if ok else 0.0,
        needs_manual_review=not ok,
        source="native",
        kind="searchable" if ok else "empty",
        error=None if ok else text[:200],
        contains_drawing=treat_as_drawing or suffix in {".dwg", ".dxf", ".rvt", ".rfa", ".rte", ".rft"},
    )
    return DocumentExtractionResult(
        pages=[page],
        merged_text=text if ok else text,
        pdf_kind="n/a",
        page_count=1,
        ocr_page_count=0,
        failed_pages=[] if ok else [1],
        needs_manual_review=not ok,
        provider=method or "legacy",
        phase=ExtractionPhase.COMPLETED.value if ok else ExtractionPhase.FAILED.value,
        extraction_method=method,
        confidence_score=82.0 if ok and method == "dxf_native" else (50.0 if ok and method else None),
    )


def has_usable_text(text: str | None) -> bool:
    if not text:
        return False
    stripped = text.strip()
    if stripped.startswith(
        (
            "[EXTRACT_ERROR]",
            "[BINARY_OR_IMAGE_FILE]",
            "[OCR_UNAVAILABLE]",
            "[OCR_EMPTY]",
            "[OCR_SKIPPED]",
            "[CAD_EMPTY]",
            "[CAD_UNSUPPORTED]",
            "[IFC_EMPTY]",
            "[IFC_ERROR]",
            "[GAEB_EMPTY]",
            "[GAEB_ERROR]",
        )
    ):
        return False
    cleaned = strip_extraction_chrome(stripped)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return len(cleaned) >= _MIN_NATIVE_TEXT_CHARS


_PAGE_HEADER_RE = re.compile(
    r"^---\s*page\s+\d+\s*\([^)]*\)\s*---(?:\s*\[[^\]]*\])?\s*",
    re.IGNORECASE | re.MULTILINE,
)
_OCR_CHROME_RE = re.compile(
    r"(?:^|\n)\s*---\s*(?:page|sheet|ocr|cad|ifc|gaeb)[^\n]*---\s*(?:\[[^\]]*\])?\s*",
    re.IGNORECASE,
)
_PERSIAN_LETTER_RE = re.compile(r"[\u0600-\u06FF]")
_KNOWN_REVERSED_FA = {
    "نامتخاس": "ساختمان",
    "هرادا": "اداره",
    "دحاو": "واحد",
    "تاملزلا": "الزامات",
    "اراکنامپ": "پیمانکارا",
    "ناراکنامپ": "پیمانکاران",
    "یمنیا": "ایمنی",
    "شراتش": "شرایط",
}
_COMMON_FA_TOKENS = {
    "از",
    "به",
    "که",
    "این",
    "را",
    "با",
    "در",
    "برای",
    "و",
    "یا",
    "تا",
    "هر",
    "یک",
    "می",
    "نمی",
    "است",
    "هست",
    "شود",
    "کل",
    "پیمانکار",
    "پیمانکاران",
    "پیمانکارا",
    "ساختمان",
    "اداره",
    "واحد",
    "ایمنی",
    "کار",
    "پروژه",
    "قرارداد",
    "الزامات",
    "شرایط",
    "محدوده",
    "شرح",
    "تاسیسات",
    "وتاسیسات",
    "دانش",
    "فنی",
    "نقشه",
    "نقشهها",
    "متره",
    "مناقصه",
    "امتیاز",
    "پیمان",
    "کارفرما",
    "نظارت",
    "مکانیک",
    "ابنیه",
}


def strip_extraction_chrome(text: str) -> str:
    """Remove OCR/page markers that must never appear in user-facing quotes."""
    cleaned = text or ""
    cleaned = re.sub(r"^\[OCR_APPLIED\][^\n]*\n?", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^---\s*cad:[^\n]*\n?", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^---\s*ifc:[^\n]*\n?", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^---\s*gaeb:[^\n]*\n?", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"^\[(?:DWG|DXF|Revit|DXF-fallback) strings\]\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = _PAGE_HEADER_RE.sub("", cleaned)
    cleaned = _OCR_CHROME_RE.sub("\n", cleaned)
    cleaned = re.sub(r"---\s*(page|sheet|ocr)[^-\n]*---", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\[OCR_TRUNCATED\][^\n]*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\[NEEDS_MANUAL_REVIEW\]", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\[(?:STAMP|SIGNATURE|DRAWING)(?:,[^\]]*)?\]", "", cleaned, flags=re.IGNORECASE)
    return cleaned.strip()


def _fa_token_score(text: str) -> int:
    toks = re.findall(r"[\u0600-\u06FF]{1,}", text or "")
    score = 0
    for t in toks:
        if t in _COMMON_FA_TOKENS:
            score += 3
        elif t.endswith(("ها", "ان", "ات", "ین", "ون")) and len(t) >= 4:
            score += 1
        elif t.startswith(("می", "نمی", "بر", "در")) and len(t) >= 3:
            score += 1
    return score


def _reverse_visual_rtl_runs(text: str) -> str:
    """Convert visually-stored RTL PDF text to logical order.

    Some native PDFs emit Persian glyphs already in visual order (each word
    character-reversed, and often the token sequence flipped). Latin/digit runs
    stay left-to-right.
    """
    if not text or not _PERSIAN_LETTER_RE.search(text):
        return text

    # Reverse whole string, then flip contiguous LTR (Latin/digit/punct) runs back.
    flipped = text[::-1]

    def _restore_ltr(m: re.Match[str]) -> str:
        return m.group(0)[::-1]

    return re.sub(r"[A-Za-z0-9][A-Za-z0-9._/%+\-]*", _restore_ltr, flipped)


def _wordwise_unreverse_persian(text: str) -> str:
    """Reverse only Arabic-script words (keeps Latin and punctuation in place)."""

    def _fix_tok(m: re.Match[str]) -> str:
        return m.group(0)[::-1]

    return re.sub(r"[\u0600-\u06FF]+", _fix_tok, text)


def _token_order_unreverse(text: str) -> str:
    """Un-reverse each Arabic word, then reverse token order (visual RTL lines)."""
    toks = (text or "").split()
    fixed: list[str] = []
    for tok in toks:
        if re.fullmatch(r"[\u0600-\u06FF]+", tok):
            fixed.append(_KNOWN_REVERSED_FA.get(tok, tok[::-1]))
        else:
            fixed.append(tok)
    return " ".join(reversed(fixed))


def _looks_visually_reversed(text: str) -> bool:
    if not text:
        return False
    # Arabic presentation forms almost always mean visual PDF order
    if re.search(r"[\uFB50-\uFDFF\uFE70-\uFEFF]", text):
        return True
    if any(k in text for k in _KNOWN_REVERSED_FA):
        return True
    reversed_hits = 0
    for tok in re.findall(r"[\u0600-\u06FF]{4,}", text):
        rev = tok[::-1]
        # Palindromes (e.g. تاسیسات) must not count — they falsely flag good text.
        if rev == tok:
            continue
        if rev in _COMMON_FA_TOKENS or rev in _KNOWN_REVERSED_FA:
            reversed_hits += 1
            if reversed_hits >= 2:
                return True
    return False


def repair_reversed_persian(text: str) -> str:
    """If Persian looks character-reversed (common in some PDFs), repair it.

    Never rewrite already-readable Persian: only accept a candidate when its
    token score is strictly better than the original.
    """
    raw = (text or "").strip()
    if not raw or not _PERSIAN_LETTER_RE.search(raw):
        return raw
    base = _fa_token_score(raw)
    candidates = [
        raw,
        _token_order_unreverse(raw),
        _wordwise_unreverse_persian(raw),
        _reverse_visual_rtl_runs(raw),
        _wordwise_unreverse_persian(_reverse_visual_rtl_runs(raw)),
    ]
    best = raw
    best_score = base
    for c in candidates:
        sc = _fa_token_score(c)
        if sc > best_score:
            best = c
            best_score = sc
    # Only rewrite when a candidate is clearly better (or known visual-RTL markers).
    if best is not raw and best_score > base:
        parts = [_KNOWN_REVERSED_FA.get(tok, tok) for tok in best.split()]
        return " ".join(parts)
    if _looks_visually_reversed(raw) and best_score >= base + 2:
        parts = [_KNOWN_REVERSED_FA.get(tok, tok) for tok in best.split()]
        return " ".join(parts)
    return raw


def clean_display_excerpt(text: str | None, *, max_len: int = 420) -> str | None:
    """Prepare a user-facing document quote: strip OCR chrome + fix RTL garble.

    Keeps the same sentence content from the file; only normalizes presentation
    glyphs and repairs *actually* reversed Persian so the quote is readable.
    """
    if not text:
        return None
    cleaned = strip_extraction_chrome(text)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if not cleaned:
        return None
    # Presentation-form glyphs (common in Iranian PDFs) → standard Arabic letters
    try:
        import unicodedata

        cleaned = unicodedata.normalize("NFKC", cleaned)
    except Exception:  # noqa: BLE001
        pass
    cleaned = repair_reversed_persian(cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    # Drop leftover page markers inside the middle of a quote
    cleaned = re.sub(r"---\s*page\s+\d+[^-\n]*---", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if len(cleaned) < 8:
        return None
    if len(cleaned) > max_len:
        cleaned = cleaned[: max_len - 1].rstrip() + "…"
    return cleaned


def verbatim_document_excerpt(text: str | None, *, max_len: int = 500) -> str | None:
    """Return a readable document quote for «جمله ایراددار».

    Strips extraction chrome and repairs only truly reversed/presentation-form
    Persian — does not invent new wording or reverse already-correct text.
    """
    return clean_display_excerpt(text, max_len=max_len)


def sentence_around_index(text: str, index: int, *, max_len: int = 500) -> str | None:
    """Cut the nearest sentence/line around ``index`` from raw extracted text.

    ``index`` must refer to the original ``text`` (not a chrome-stripped copy),
    so the returned slice matches the document characters exactly.
    """
    if not text:
        return None
    raw = text
    if not raw.strip():
        return None
    idx = max(0, min(int(index), len(raw) - 1))
    # Prefer paragraph/line boundaries, then sentence punctuation.
    start = 0
    for sep in ("\n", "؟", "!", "?", ".", "。"):
        pos = raw.rfind(sep, 0, idx)
        if pos >= 0:
            start = max(start, pos + 1)
            break
    else:
        start = max(0, idx - 120)

    end = len(raw)
    for sep in ("\n", "؟", "!", "?", ".", "。"):
        pos = raw.find(sep, idx)
        if pos >= 0:
            end = min(end, pos + (0 if sep == "\n" else 1))
            break
    else:
        end = min(len(raw), idx + 280)

    chunk = raw[start:end].strip()
    if len(chunk) < 12:
        # Fallback window around the hit
        chunk = raw[max(0, idx - 80) : min(len(raw), idx + 220)].strip()
    return clean_display_excerpt(chunk, max_len=max_len)


# Same threshold used by STD-SEED-005 / standard_extraction_confidence.
EXTRACTION_CONFIDENCE_LIMITATION_THRESHOLD = 50.0


def _coerce_confidence(val: object) -> float | None:
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def get_extraction_confidence(doc: object) -> float | None:
    """Read extraction confidence from DocView, document dict, or meta blobs.

    Missing confidence stays None — never invent 0 (that falsely flags limitations).
    Does not read a `.confidence` property (avoids recursion with DocView).
    """
    if doc is None:
        return None

    meta: dict = {}
    canonical: dict = {}
    top: object = None

    if isinstance(doc, dict):
        top = doc.get("confidence_score", doc.get("confidence"))
        raw_meta = doc.get("meta_json", doc.get("meta"))
        if isinstance(raw_meta, dict):
            meta = raw_meta
        elif isinstance(raw_meta, str) and raw_meta.strip():
            try:
                parsed = json.loads(raw_meta)
                if isinstance(parsed, dict):
                    meta = parsed
            except (json.JSONDecodeError, TypeError):
                meta = {}
        can = doc.get("canonical")
        if isinstance(can, dict):
            canonical = can
        elif isinstance(meta.get("canonical"), dict):
            canonical = meta["canonical"]
    else:
        top = getattr(doc, "confidence_score", None)
        raw_meta = getattr(doc, "meta", None)
        if raw_meta is None:
            raw_meta = getattr(doc, "meta_json", None)
        if isinstance(raw_meta, dict):
            meta = raw_meta
        elif isinstance(raw_meta, str) and raw_meta.strip():
            try:
                parsed = json.loads(raw_meta)
                if isinstance(parsed, dict):
                    meta = parsed
            except (json.JSONDecodeError, TypeError):
                meta = {}
        can = getattr(doc, "canonical", None)
        if isinstance(can, dict):
            canonical = can
        elif isinstance(meta.get("canonical"), dict):
            canonical = meta["canonical"]

    conf = _coerce_confidence(top)
    if conf is not None:
        return conf

    for blob in (canonical.get("extraction"), meta.get("extraction"), meta):
        if not isinstance(blob, dict):
            continue
        conf = _coerce_confidence(blob.get("confidence_score") or blob.get("confidenceScore"))
        if conf is not None:
            return conf
    return None


def has_low_extraction_confidence(doc: object) -> bool:
    """True when confidence is known and below the shared limitation threshold."""
    conf = get_extraction_confidence(doc)
    return conf is not None and conf < EXTRACTION_CONFIDENCE_LIMITATION_THRESHOLD


def document_has_extraction_limitation(
    doc: object,
    *,
    category: object | None = None,
    extracted_text: str | None = None,
) -> bool:
    """Unified limitation flag: unusable text (non-drawings) OR low extraction confidence.

    Empty drawings do not count as limitations (they are excluded from the text gate).
    """
    from app.models import DocumentCategory

    if category is None:
        category = getattr(doc, "category", None) if not isinstance(doc, dict) else doc.get("category")
    cat_val = category.value if isinstance(category, DocumentCategory) else str(category or "")

    if extracted_text is None:
        if isinstance(doc, dict):
            extracted_text = doc.get("extracted_text") or ""
        else:
            extracted_text = getattr(doc, "extracted_text", None) or ""

    is_drawing = cat_val == DocumentCategory.DRAWING.value
    if not is_drawing and not has_usable_text(extracted_text):
        return True
    return has_low_extraction_confidence(doc)


def ocr_status() -> dict:
    """Runtime OCR capability for health/debug — always defined (safe import)."""
    try:
        return ocr_runtime_status()
    except Exception as exc:  # noqa: BLE001
        return {"available": False, "error": str(exc)}


def _extract_image(path: Path) -> str:
    from app.services.ocr.factory import get_ocr_service

    svc = get_ocr_service()
    if svc is None:
        return (
            f"[OCR_UNAVAILABLE] {path.name}\n"
            "OCR engine not installed."
        )
    from PIL import Image

    with Image.open(path) as img:
        result = svc.recognize_image(img.copy(), page_number=1)
    if result.text.strip():
        return f"--- ocr: {path.name} ---\n{result.text.strip()}"
    return f"[OCR_EMPTY] {path.name}\nImage OCR produced no usable text."


def _extract_docx(path: Path, *, allow_ocr: bool = True) -> str:
    doc = DocxDocument(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    native = "\n".join(parts).strip()
    if len(re.sub(r"\s+", "", native)) >= _MIN_NATIVE_TEXT_CHARS:
        return native

    if not allow_ocr:
        return native or (
            f"[OCR_SKIPPED] {path.name}\n"
            "DOCX had little/no text; OCR of embedded images was skipped."
        )

    image_texts = _ocr_docx_images(doc)
    combined = "\n\n".join([p for p in [native, image_texts] if p]).strip()
    if len(re.sub(r"\s+", "", combined)) >= _MIN_NATIVE_TEXT_CHARS:
        return "[OCR_APPLIED] DOCX images OCR'd.\n\n" + combined
    return combined or (
        f"[OCR_EMPTY] {path.name}\nWord file had no readable text or OCR-able images."
    )


def _ocr_docx_images(doc: DocxDocument) -> str:
    from app.services.ocr.factory import get_ocr_service
    import io
    from PIL import Image

    svc = get_ocr_service()
    if svc is None:
        return ""
    chunks: list[str] = []
    count = 0
    try:
        for rel in doc.part.rels.values():
            if count >= _MAX_DOCX_IMAGES:
                break
            rel_type = getattr(rel, "reltype", "") or ""
            if "image" not in rel_type:
                continue
            try:
                blob = rel.target_part.blob
            except Exception:  # noqa: BLE001
                continue
            count += 1
            try:
                with Image.open(io.BytesIO(blob)) as img:
                    result = svc.recognize_image(img.copy(), page_number=count)
                if len(result.text.strip()) >= _MIN_PAGE_TEXT_CHARS:
                    chunks.append(f"--- ocr image {count} ---\n{result.text.strip()}")
            except Exception as exc:  # noqa: BLE001
                logger.warning("DOCX image OCR failed: %s", exc)
    except Exception as exc:  # noqa: BLE001
        logger.warning("DOCX image walk failed: %s", exc)
    return "\n\n".join(chunks).strip()


def _extract_excel(path: Path) -> str:
    wb = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"--- sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if values:
                parts.append(" | ".join(values))
    return "\n".join(parts)


def merge_extraction_meta(existing_meta: str | None, extraction_meta: dict) -> str:
    """Shallow merge at top level; deep-merge `extraction` so CDM/canonical is not wiped."""
    base: dict = {}
    if existing_meta:
        try:
            parsed = json.loads(existing_meta)
            if isinstance(parsed, dict):
                base = parsed
        except json.JSONDecodeError:
            base = {}
    for key, value in (extraction_meta or {}).items():
        if (
            key == "extraction"
            and isinstance(value, dict)
            and isinstance(base.get("extraction"), dict)
        ):
            merged = dict(base["extraction"])
            merged.update(value)
            base["extraction"] = merged
        else:
            base[key] = value
    return json.dumps(base, ensure_ascii=False)
